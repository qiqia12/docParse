from io import BytesIO
from typing import Iterator

from pptx import Presentation

from ..base_parser import BaseParser
from ..models import ParseChunk, ParseOptions


class PptxParser(BaseParser):
    def supported_formats(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ]

    def parse(self, file_bytes: bytes, filename: str, options: ParseOptions) -> Iterator[ParseChunk]:
        prs = Presentation(BytesIO(file_bytes))
        total_slides = len(prs.slides)

        if total_slides == 0:
            yield ParseChunk(page_number=1, total_pages=1, markdown="", is_complete=True)
            return

        for slide_num, slide in enumerate(prs.slides, 1):
            lines = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            if shape.is_placeholder and shape.placeholder_format.type == 1:
                                lines.append(f"## {text}")
                            else:
                                lines.append(text)

                if shape.has_table:
                    table = shape.table
                    header = "| " + " | ".join(cell.text for cell in table.rows[0].cells) + " |"
                    sep = "|" + "|".join("---" for _ in table.rows[0].cells) + "|"
                    rows = []
                    for row in table.rows[1:]:
                        rows.append("| " + " | ".join(cell.text for cell in row.cells) + " |")
                    lines.append("\n".join([header, sep] + rows))

            yield ParseChunk(
                page_number=slide_num,
                total_pages=total_slides,
                markdown="\n\n".join(lines),
                is_complete=(slide_num == total_slides),
            )
