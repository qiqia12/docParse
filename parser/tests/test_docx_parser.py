import pytest
from src.parsers.docx_parser import DocxParser
from src.models import ParseOptions


def test_docx_parser_supported_formats():
    parser = DocxParser()
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in parser.supported_formats()


def test_docx_parser_with_bytes():
    parser = DocxParser()
    from docx import Document
    from io import BytesIO
    doc = Document()
    doc.add_paragraph("Hello World")
    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    chunks = list(parser.parse(docx_bytes, "test.docx", ParseOptions()))
    assert len(chunks) >= 1
    assert any("Hello World" in c.markdown for c in chunks)


def test_docx_parser_empty_document():
    parser = DocxParser()
    from docx import Document
    from io import BytesIO
    doc = Document()
    buf = BytesIO()
    doc.save(buf)
    chunks = list(parser.parse(buf.getvalue(), "empty.docx", ParseOptions()))
    assert len(chunks) >= 1


def test_docx_parser_invalid_bytes():
    parser = DocxParser()
    # MarkItDown gracefully handles invalid bytes, returning them as raw text
    chunks = list(parser.parse(b"not a docx file", "fake.docx", ParseOptions()))
    assert len(chunks) >= 1
