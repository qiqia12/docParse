import pytest
from src.parsers.pptx_parser import PptxParser
from src.models import ParseOptions


def test_pptx_parser_supported_formats():
    parser = PptxParser()
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in parser.supported_formats()


def test_pptx_parser_with_bytes():
    from pptx import Presentation
    from io import BytesIO

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Slide Title"
    buf = BytesIO()
    prs.save(buf)

    parser = PptxParser()
    chunks = list(parser.parse(buf.getvalue(), "test.pptx", ParseOptions()))
    assert len(chunks) >= 1
    assert any("Slide Title" in c.markdown for c in chunks)


def test_pptx_parser_page_per_slide():
    from pptx import Presentation
    from io import BytesIO

    prs = Presentation()
    for i in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = f"Slide {i+1}"

    buf = BytesIO()
    prs.save(buf)

    parser = PptxParser()
    chunks = list(parser.parse(buf.getvalue(), "multi.pptx", ParseOptions()))
    assert len(chunks) == 3
    for i, chunk in enumerate(chunks):
        assert chunk.page_number == i + 1
        assert chunk.total_pages == 3
